# -*- coding: utf-8 -*-
"""Lê 6 questões de um .docx, preenche campos reconhecidos no Brave e NÃO envia."""
import json, os, sys, time, unicodedata, re
from pathlib import Path
from docx import Document
try:
    from selenium import webdriver
except ImportError:
    webdriver=None
try:
    from selenium.webdriver.common.by import By
    from selenium.webdriver.common.keys import Keys
    from selenium.webdriver.chrome.options import Options
except ImportError:
    By=Keys=Options=None

ROOT=Path(__file__).resolve().parent
FORM_URL='https://forms.gle/pmyGEukaMMk2DFLMA'

def norm(s):
    return ' '.join(''.join(c for c in unicodedata.normalize('NFKD',str(s).lower()) if not unicodedata.combining(c)).split())

def load_doc(path, requested_cap=None):
    """Lê seções CAPÍTULO N, com opção exata do Forms especificada no Word."""
    doc=Document(path)
    meta={}; groups={}; topics={}; active=None
    for table in doc.tables:
        rows=[[c.text.strip() for c in r.cells] for r in table.rows]
        if not rows:continue
        title=rows[0][0].strip()
        match=re.match(r'^CAP[IÍ]TULO\s+(\d+)\b',title,re.I)
        if match:
            active=match.group(1)
            if active in groups:raise ValueError(f'Capítulo {active} repetido. Use um número diferente em cada seção.')
            groups[active]=[]
            for r in rows[1:]:
                if len(r)>1 and norm(r[0]) in ('topico no formulario','topico da questao no formulario'):
                    topics[active]=r[1].strip()
        elif norm(title).startswith('questao ' ) or (norm(title)=='questao' and len(rows)>=4):
            if active is None:raise ValueError('Há uma questão antes do cabeçalho CAPÍTULO N.')
            q={norm(r[0]):r[1].strip() for r in (rows if norm(title)=='questao' else rows[1:]) if len(r)>=2}
            groups[active].append(q)
        else:
            meta.update({norm(r[0]):r[1].strip() for r in rows if len(r)>=2})
    for k in ['nome completo','matricula','e-mail']:
        if not meta.get(k) or meta[k].startswith('['):raise ValueError(f'Preencha {k} na identificação.')
    if not groups:raise ValueError('Nenhum cabeçalho CAPÍTULO N encontrado.')
    if requested_cap is None:
        if len(groups)>1:raise ValueError('Informe o número do capítulo desejado com --cap N, ou use INICIAR.bat.')
        cap=next(iter(groups))
    else:cap=str(requested_cap)
    if cap not in groups:raise ValueError(f'Capítulo {cap} não foi encontrado. Disponíveis: {", ".join(groups)}')
    topic=topics.get(cap,'')
    if not topic or topic.startswith('['):raise ValueError(f'Capítulo {cap}: preencha a opção exata do campo Tópico no formulário.')
    questions=groups[cap]
    if len(questions)!=6:raise ValueError(f'Capítulo {cap}: esperadas 6 questões, encontradas {len(questions)}.')
    data=[]
    for i,q in enumerate(questions,1):
        for k in ['questao','resposta','citacoes e referencias','pertinencia']:
            if not q.get(k) or q[k].startswith('['):raise ValueError(f'Capítulo {cap}, questão {i}: preencha {k}.')
        r=norm(q['resposta'])
        if r not in ['verdadeira','falsa','verdadeiro','falso']:raise ValueError(f'Capítulo {cap}, questão {i}: resposta deve ser Verdadeira ou Falsa.')
        data.append({'topico':topic,'questao':q['questao'],'resposta':'Verdadeira' if r.startswith('verdadeir') else 'Falsa','citacao':q['citacoes e referencias'],'pertinencia':q['pertinencia'],'nome':meta['nome completo'],'matricula':meta['matricula'],'email':meta['e-mail']})
    return cap,data

def list_chapters(path):
    doc=Document(path)
    found=[]
    for table in doc.tables:
        if not table.rows:continue
        title=table.cell(0,0).text.strip()
        match=re.match(r'^CAP[IÍ]TULO\s+(\d+)\b',title,re.I)
        if match:
            topic=next((row.cells[1].text.strip() for row in table.rows[1:] if len(row.cells)>1 and norm(row.cells[0].text)=='topico no formulario'),'')
            found.append((match.group(1),topic))
    return found

ALIASES={
 'topico':['topico da questao','topico','assunto','tema da questao'],
 'questao':['enunciado da questao','questao','afirmacao','pergunta'],
 'resposta':['resposta correta','gabarito','verdadeiro ou falso','resposta'],
 'citacao':['citacoes e referencias','citacao e referencia','referencias','citacao','fonte bibliografica'],
 'pertinencia':['pertinencia','justificativa','relevancia'],
 'nome':['nome completo','seu nome','nome do aluno'],
 'matricula':['matricula','numero de matricula'],
 'email':['email','e-mail','endereco de e-mail']}

def label(block):
    hs=block.find_elements(By.CSS_SELECTOR,'[role="heading"], .M7eMe, .HoXoMd')
    return (hs[0].text if hs and hs[0].text.strip() else (block.text or '').split('\n')[0]).strip()

def identify(s):
    r=norm(s)
    for k,aliases in ALIASES.items():
        if any(a==r or r.startswith(a+' ') or r.startswith(a+':') for a in aliases):return k
    for k,aliases in ALIASES.items():
        if any(a in r for a in aliases):return k
    return None

def choose(driver, block, desired):
    wanted=norm(desired)
    def matches(el):
        texts=[el.get_attribute('aria-label') or '',el.get_attribute('data-value') or '',el.text or '']
        return any(norm(x).strip(' *')==wanted or (wanted in ['verdadeira','falsa'] and norm(x) in [wanted[:-1]+'o',wanted]) for x in texts)
    for el in block.find_elements(By.CSS_SELECTOR,'[role="radio"], input[type="radio"], [role="option"]'):
        if not el.is_displayed():continue
        candidates=[el]
        try:candidates += [el.find_element(By.XPATH,'..'),el.find_element(By.XPATH,'../..')]
        except Exception:pass
        if any(matches(x) for x in candidates):
            driver.execute_script('arguments[0].scrollIntoView({block:"center"});',el)
            try:el.click()
            except Exception:driver.execute_script('arguments[0].click()',el)
            time.sleep(.12)
            if el.get_attribute('aria-checked')=='true' or el.is_selected():return True
    for el in block.find_elements(By.CSS_SELECTOR,'[data-value], [aria-label], label'):
        if not el.is_displayed() or not matches(el):continue
        try:
            el.click()
            if block.find_elements(By.CSS_SELECTOR,'[role="radio"][aria-checked="true"],input[type="radio"]:checked'):return True
        except Exception:pass
    for box in block.find_elements(By.CSS_SELECTOR,'[role="listbox"]'):
        if not box.is_displayed():continue
        box.click();time.sleep(.3)
        for el in driver.find_elements(By.CSS_SELECTOR,'[role="option"]'):
            if el.is_displayed() and matches(el):el.click();return True
        driver.switch_to.active_element.send_keys(Keys.ESCAPE)
    return False

def fill(driver,item,number):
    driver.get(FORM_URL);time.sleep(1.8)
    found=[];failed=[]
    for block in driver.find_elements(By.CSS_SELECTOR,'[role="listitem"]'):
        if not block.is_displayed():continue
        name=label(block); key=identify(name)
        if key not in item:continue
        try:
            ok=False
            if key!='topico':
                edits=[el for el in block.find_elements(By.CSS_SELECTOR,'input[type="text"],input[type="email"],textarea,[contenteditable="true"]') if el.is_displayed()]
                if edits:
                    edits[0].click();edits[0].send_keys(Keys.CONTROL,'a');edits[0].send_keys(item[key]);ok=True
            if not ok:ok=choose(driver,block,item[key])
            (found if ok else failed).append(key)
        except Exception as ex:failed.append(f'{key}:{type(ex).__name__}')
    print(f'Questão {number}: preenchidos: {", ".join(found) or "nenhum"}; pendentes: {", ".join(failed) or "nenhum"}',flush=True)
    if 'topico' not in found:print('   ATENÇÃO: CONFIRA A BOLINHA DO TÓPICO!',flush=True)
    if not found:print('   Confira login, campos ou mudanças no Google Forms.',flush=True)


def main():
    if len(sys.argv)<2:
        print('Uso: py -3 preencher_brave.py "documento.docx" --listar | --cap NUMERO [--validar]')
        return 2
    path=Path(sys.argv[1].strip('"')).expanduser()
    if not path.is_file():raise ValueError(f'Arquivo não encontrado: {path}')
    if '--listar' in sys.argv:
        for chapter,topic in list_chapters(path):print(f'Capítulo {chapter} — opção do Forms: {topic}')
        return 0
    cap_arg=next((sys.argv[i+1] for i,v in enumerate(sys.argv[:-1]) if v=='--cap'),None)
    cap,items=load_doc(path,cap_arg)
    print(f'Validado: capítulo {cap}, 6 questões; aluno: {items[0]["nome"]}.')
    print(f'Tópico para clicar no Forms: {items[0]["topico"]}')
    if '--validar' in sys.argv:
        print('Estrutura válida. Isso não comprova o conteúdo factual ou a disponibilidade dessa opção no Forms.')
        return 0
    if webdriver is None:raise RuntimeError('Selenium não instalado. Execute INSTALAR.bat.')
    opts=Options();opts.add_experimental_option('detach',True);opts.add_argument('--start-maximized')
    paths=[os.environ.get('BRAVE_PATH',''),*[os.path.join(os.environ.get(k,''),'BraveSoftware','Brave-Browser','Application','brave.exe') for k in ['PROGRAMFILES','PROGRAMFILES(X86)','LOCALAPPDATA']]]
    brave=next((x for x in paths if x and Path(x).is_file()),None)
    if not brave:raise RuntimeError('Brave não encontrado. Instale Brave ou configure BRAVE_PATH.')
    opts.binary_location=brave
    browser=webdriver.Chrome(options=opts)
    print('Abrindo seis abas para conferência, sem enviar respostas.')
    for i,item in enumerate(items,1):
        if i>1:browser.switch_to.new_window('tab')
        try:fill(browser,item,i)
        except Exception as ex:print(f'Questão {i}: erro: {ex}')
    print('Confira TODOS os campos e envie manualmente. Se o tópico não existe no formulário, ele NÃO será marcado.')
    return 0
if __name__=='__main__':
    try:sys.exit(main())
    except Exception as ex:print('ERRO:',ex);sys.exit(1)
